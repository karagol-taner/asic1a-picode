#!/usr/bin/env python3
"""
asic_fixed_effects.py - fixed-effects model of the measured pH50 shift.

Usage:
    python3 asic_fixed_effects.py asic1a_ph50.csv \
        --props 42003_2025_9090_MOESM4_ESM.xlsx [--readout act|ssd]
        [--edges core|implemented]

  |delta pH50| ~ position fixed effect + |delta charge| + |delta volume| + ladder

Position fixed effects are applied by within-position centering. Positions
carrying a single substitution are dropped, and the count is reported.
Significance for the ladder term is by 20,000 shuffles of the indicator within
position; the t statistic is printed for reference. Side-chain volumes come
from the source study's Supplementary Data 1.

Exits non-zero on any fatal condition.
"""
import argparse, csv, os, re, sys
import numpy as np

RNG_SEED = 20261007
N_PERM = 20000


def die(msg):
    sys.exit('FATAL asic_fixed_effects.py: %s' % msg)


def load_ladder(path=None):
    """The ladder edges and charges, by way of codes/picode_paths.py.

    An explicit --ladder path wins. Otherwise the pi-code method folder is used
    where it is present and the snapshot committed in data/ where it is not, so
    the definitions keep a single home and a clone of this repository without
    the method folder still resolves the same published constants.
    """
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import picode_paths as P
    return P.load_ladder(path=path)


def harvest_props(xlsx):
    """Per-residue side-chain properties, read from the source workbook."""
    try:
        import openpyxl
    except ImportError:
        die('openpyxl is not installed')
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    props, header = {}, None
    for sh in wb.sheetnames:
        if not re.match(r'^[A-Z]\d{1,3}_(act|des)$', sh):
            continue
        for i, row in enumerate(wb[sh].iter_rows(values_only=True)):
            if i == 0:
                header = [str(c) if c else '' for c in row]
                continue
            if not row or not isinstance(row[0], str) or len(row[0]) != 1:
                continue
            if not row[0].isupper():
                continue                       # the trailing 'n' count row
            vals = row[2:12]
            if all(isinstance(v, (int, float)) for v in vals):
                props.setdefault(row[0], [float(v) for v in vals])
    if not props or header is None:
        die('no per-residue property rows harvested from %s' % xlsx)
    if header[2] != 'size':
        die('expected the third column of the property block to be "size", found %r'
            % header[2])
    return props, header[2:12]


def ols(y, X):
    """Least squares with HC0 robust standard errors."""
    XtX = X.T @ X
    if np.linalg.matrix_rank(XtX) < X.shape[1]:
        die('design matrix is rank deficient; a predictor is collinear with another')
    XtXi = np.linalg.inv(XtX)
    b = XtXi @ (X.T @ y)
    r = y - X @ b
    S = (X * r[:, None]).T @ (X * r[:, None])
    cov = XtXi @ S @ XtXi
    se = np.sqrt(np.diag(cov))
    return b, se, r


def demean(y, X, groups):
    """Within-group centering. Returns centred arrays and the kept mask."""
    y2, X2 = y.astype(float).copy(), X.astype(float).copy()
    keep = np.zeros(len(y), bool)
    for g in np.unique(groups):
        m = groups == g
        if m.sum() < 2:
            continue
        keep |= m
        y2[m] -= y2[m].mean()
        X2[m] -= X2[m].mean(axis=0)
    return y2[keep], X2[keep], keep


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('csv')
    ap.add_argument('--ladder', help='path to ladder_edge_cost.py; by default the method folder, then data/ladder_definition.json')
    ap.add_argument('--props', required=True)
    ap.add_argument('--readout', choices=['act', 'ssd'], default='act')
    ap.add_argument('--edges', choices=['core', 'implemented'], default='core',
                    help='core = the 12 first-position edges the method defines, '
                         'which is the primary analysis; implemented = the 17 the '
                         'tool routes (12 core + 5 arginine), reported in the '
                         'manuscript as a sensitivity')
    a = ap.parse_args()

    L = load_ladder(a.ladder)
    props, pnames = harvest_props(a.props)
    size_i = pnames.index('size')
    edges = set(L.CORE) | (set(L.ARG) if a.edges == 'implemented' else set())

    col = 'd_act' if a.readout == 'act' else 'd_ssd'
    label = 'activation pH50' if a.readout == 'act' else 'SSD pHD50'
    rows = [r for r in csv.DictReader(open(a.csv))
            if r['mutation'] != 'WT' and r.get(col) not in (None, '')]
    if not rows:
        die('no usable substitution rows in %s' % a.csv)

    miss = sorted({r[k] for r in rows for k in ('wt_aa', 'mut_aa')} - set(props))
    if miss:
        die('no harvested side-chain properties for %s; the model would silently '
            'drop those rows' % ', '.join(miss))

    y, dq, dv, lad, pos, name = [], [], [], [], [], []
    for r in rows:
        w, m = r['wt_aa'], r['mut_aa']
        y.append(abs(float(r[col])))
        dq.append(abs(L.CHARGE.get(m, 0.0) - L.CHARGE.get(w, 0.0)))
        dv.append(abs(props[m][size_i] - props[w][size_i]))
        lad.append(1.0 if (w, m) in edges else 0.0)
        pos.append(int(r['resi']))
        name.append(r['mutation'])
    y = np.array(y); pos = np.array(pos); lad = np.array(lad)
    X = np.column_stack([dq, dv, lad])
    names = ['|delta charge|', '|delta volume|', 'ladder']

    print('readout: %s      edge set: %s (%d directed edges)'
          % (label, a.edges, len(edges)))
    print('substitutions: %d over %d positions; on a ladder edge: %d'
          % (len(y), len(set(pos)), int(lad.sum())))

    yc, Xc, keep = demean(y, X, pos)
    used_pos = sorted(set(pos[keep]))
    dropped = sorted(set(pos) - set(used_pos))
    nlad = int(lad[keep].sum())
    print('after within-position centering: %d substitutions over %d positions '
          'carrying 2 or more' % (len(yc), len(used_pos)))
    print('dropped as singletons: %d positions %s' % (len(dropped), dropped))
    if nlad == 0 or nlad == int(keep.sum()):
        die('the ladder indicator has no within-position variation after centering, '
            'so its coefficient is not identified')

    b, se, resid = ols(yc, Xc)
    df = len(yc) - len(used_pos) - Xc.shape[1]
    if df <= 0:
        die('no residual degrees of freedom: %d observations, %d positions, %d '
            'predictors' % (len(yc), len(used_pos), Xc.shape[1]))
    print('residual degrees of freedom: %d\n' % df)

    print('%-16s %-11s %-11s %-9s' % ('term', 'coef (pH)', 'robust SE', 't'))
    for i, nm in enumerate(names):
        print('%-16s %+-11.4f %-11.4f %+-9.2f' % (nm, b[i], se[i], b[i] / se[i]))

    # Permutation: shuffle ladder membership within position.
    rng = np.random.default_rng(RNG_SEED)
    li = names.index('ladder')
    obs = b[li]
    pg = pos[keep]
    lad_k = Xc[:, li].copy()
    base = Xc.copy()
    hits = 0
    perm_lad = lad[keep].copy()
    for _ in range(N_PERM):
        sh = perm_lad.copy()
        for g in np.unique(pg):
            m = pg == g
            sh[m] = rng.permutation(sh[m])
        col_sh = sh.astype(float).copy()
        for g in np.unique(pg):
            m = pg == g
            col_sh[m] -= col_sh[m].mean()
        if np.allclose(col_sh, 0):
            continue
        Xp = base.copy()
        Xp[:, li] = col_sh
        bp, _, _ = ols(yc, Xp)
        if abs(bp[li]) >= abs(obs) - 1e-12:
            hits += 1
    pperm = (hits + 1) / float(N_PERM + 1)
    print('\nladder coefficient %+.4f pH units' % obs)
    print('permutation p = %.4f  (%d shuffles of ladder membership within position)'
          % (pperm, N_PERM))
    print('interpretation: a ladder substitution perturbs %s by %.3f pH units %s '
          'than a non-ladder substitution at the same position with the same charge '
          'change and a comparable volume change.'
          % (label, abs(obs), 'LESS' if obs < 0 else 'MORE'))

    # Charge-matched subset, the most conservative form of the same comparison.
    print('\n--- restricted to positions holding both a ladder and a non-ladder '
          'substitution at the SAME charge change ---')
    key = {}
    for i in np.where(keep)[0]:
        key.setdefault((pos[i], round(X[i, 0], 3)), []).append(i)
    sel = [i for v in key.values() for i in v
           if len({lad[j] for j in v}) == 2]
    if len(sel) < 8:
        print('  only %d substitutions qualify; not reported' % len(sel))
        return
    ys, Xs, ks = demean(y[sel], X[sel], pos[sel])
    bs, ses, _ = ols(ys, Xs)
    print('  %d substitutions over %d positions; ladder coefficient %+.4f pH '
          '(robust SE %.4f)' % (len(ys), len(set(pos[sel][ks])), bs[li], ses[li]))


if __name__ == '__main__':
    main()
